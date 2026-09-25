from docx import Document
from pathlib import Path
for p in [Path(r'D:\work\2026\9\1 集团党组巡视河南整改台账-焦作市分公司\集团党组巡视整改落实方案-2026年9月6日(1).docx'),Path(r'D:\work\2026\9\2 巡察整改发言\发言.docx'),Path(r'D:\work\2026\9\2 杜蘅自查整改\专项整改报告确定版.docx')]:
 print('====',p)
 try:
  d=Document(str(p))
  for para in d.paragraphs:
   t=para.text.strip()
   if any(k in t for k in ['9月','9 月','整改报告','报送','完成时限','整改台账']): print(t)
  for table in d.tables:
   for row in table.rows:
    txt=' | '.join(c.text.strip().replace('\n',' ') for c in row.cells)
    if any(k in txt for k in ['9月','9 月','整改报告','报送','完成时限']): print(txt)
 except Exception as e: print(e)
