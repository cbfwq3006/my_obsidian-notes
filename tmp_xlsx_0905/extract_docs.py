from pathlib import Path
from docx import Document
import json
root=Path(r'D:\work\2026\9\1 集团党组巡视河南整改台账-焦作市分公司')
out={}
for p in root.rglob('*.docx'):
    if p.name.startswith('~$'): continue
    try:
        d=Document(str(p))
        parts=[]
        for para in d.paragraphs:
            if para.text.strip(): parts.append(para.text.strip())
        for tbl in d.tables:
            for row in tbl.rows:
                parts.append(' | '.join(cell.text.replace('\\n',' / ').strip() for cell in row.cells))
        out[str(p)]=parts
    except Exception as e:
        out[str(p)]=[f'ERROR {e}']
Path(r'E:\sl_obsidian\tmp_xlsx_0905\docs.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print('docs',len(out))
for k,v in out.items(): print(k, len(v), sum(map(len,v)))
