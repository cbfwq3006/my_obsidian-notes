from docx import Document
from pathlib import Path
import json,sys
root=Path(r'D:\work\2026\9\3 王统主任安排工作\自查报告（7项）')
out=[]
for p in sorted(root.glob('*.docx')):
 d=Document(str(p)); paras=[x.text for x in d.paragraphs if x.text.strip()]
 tables=[]
 for t in d.tables:
  for row in t.rows:
   tables.append([c.text for c in row.cells])
 out.append({'file':p.name,'paras':paras,'tables':tables})
Path(r'E:\sl_obsidian\tmp_docs.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
