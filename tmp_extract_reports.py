from pathlib import Path
from docx import Document
base=Path(r"D:\work\2026\9\3 王统主任安排工作\自查报告（7项）")
for f in sorted(base.glob('*.docx')):
 print('\n###',f.name)
 d=Document(f)
 for p in d.paragraphs:
  t=' '.join(p.text.split())
  if t: print(t)
 for ti,tbl in enumerate(d.tables):
  print(f'--TABLE {ti}--')
  for row in tbl.rows:
   print(' | '.join(' '.join(c.text.split()) for c in row.cells))
