from docx import Document
import json,sys
p=sys.argv[1]; outp=sys.argv[2]
doc=Document(p)
out={'paragraphs':[], 'tables':[]}
for i,para in enumerate(doc.paragraphs):
    if para.text.strip(): out['paragraphs'].append({'i':i,'style':para.style.name if para.style else None,'text':para.text})
for ti,t in enumerate(doc.tables):
    rows=[]
    for ri,row in enumerate(t.rows):
        cells=[]
        for ci,c in enumerate(row.cells): cells.append({'text':c.text,'paras':[p.text for p in c.paragraphs]})
        rows.append(cells)
    out['tables'].append({'i':ti,'rows':rows,'nrows':len(rows),'ncols':len(t.columns)})
with open(outp,'w',encoding='utf-8') as f: json.dump(out,f,ensure_ascii=False,indent=2)
