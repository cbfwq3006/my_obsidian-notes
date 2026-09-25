from pathlib import Path
from docx import Document
import sys
sys.stdout.reconfigure(encoding='utf-8')
base=Path(r"D:\work\2026\9\3 王统主任安排工作\自查报告（7项）")
out=[]
for f in sorted(base.glob('*.docx')):
 out.append('\n### '+f.name)
 d=Document(f)
 for p in d.paragraphs:
  t=' '.join(p.text.split())
  if t: out.append(t)
 for ti,tbl in enumerate(d.tables):
  out.append(f'--TABLE {ti}--')
  for row in tbl.rows:
   out.append(' | '.join(' '.join(c.text.split()) for c in row.cells))
print('\n'.join(out))
